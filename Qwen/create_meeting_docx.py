
import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tcPr.append(shd)


def set_cell_text_bold(cell):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)


def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr

    borders = tblPr.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):

        tag = "w:" + edge

        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFBFBF")


def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)

    if level == 1:
        heading.runs[0].font.size = Pt(18)

    elif level == 2:
        heading.runs[0].font.size = Pt(13)

    return heading


def add_bullets(doc, lines):

    for line in lines:

        line = line.strip()

        if not line:
            continue

        line = re.sub(r"^[-*]\s*", "", line)
        line = re.sub(r"^\d+\.\s*", "", line)

        paragraph = doc.add_paragraph(
            style="List Bullet"
        )

        paragraph.add_run(line)


def parse_markdown_table(doc, lines, start_index):

    rows = []

    i = start_index

    while i < len(lines):

        line = lines[i].strip()

        if not line.startswith("|"):
            break

        # Ignore separator row
        if re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", line):
            i += 1
            continue

        cells = [
            c.strip()
            for c in line.strip("|").split("|")
        ]

        rows.append(cells)

        i += 1

    if not rows:
        return i

    columns = max(len(row) for row in rows)

    table = doc.add_table(
        rows=len(rows),
        cols=columns
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    add_table_borders(table)

    for r, row in enumerate(rows):

        for c in range(columns):

            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            if c < len(row):
                cell.text = row[c]

            if r == 0:
                set_cell_shading(cell, "D9E2F3")
                set_cell_text_bold(cell)

    doc.add_paragraph()

    return i


def add_markdown_content(doc, text):

    lines = text.splitlines()

    i = 0

    paragraph_buffer = []

    def flush_paragraph():

        nonlocal paragraph_buffer

        if paragraph_buffer:

            content = " ".join(
                x.strip()
                for x in paragraph_buffer
                if x.strip()
            )

            if content:
                doc.add_paragraph(content)

            paragraph_buffer = []

    while i < len(lines):

        line = lines[i].rstrip()

        # Blank line
        if not line.strip():

            flush_paragraph()
            i += 1
            continue

        # H1
        if line.startswith("# "):

            flush_paragraph()

            heading_text = line[2:].strip()

            heading = doc.add_heading(
                heading_text,
                level=1
            )

            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

            i += 1
            continue

        # H2
        if line.startswith("## "):

            flush_paragraph()

            heading_text = line[3:].strip()

            add_heading(
                doc,
                heading_text,
                2
            )

            i += 1
            continue

        # H3
        if line.startswith("### "):

            flush_paragraph()

            heading_text = line[4:].strip()

            add_heading(
                doc,
                heading_text,
                3
            )

            i += 1
            continue

        # Markdown table
        if (
            line.startswith("|")
            and i + 1 < len(lines)
            and "|" in lines[i + 1]
        ):

            flush_paragraph()

            i = parse_markdown_table(
                doc,
                lines,
                i
            )

            continue

        # Bullet
        if re.match(r"^\s*[-*]\s+", line):

            flush_paragraph()

            bullet_lines = []

            while i < len(lines):

                current = lines[i].strip()

                if not re.match(
                    r"^[-*]\s+",
                    current
                ):
                    break

                bullet_lines.append(current)
                i += 1

            add_bullets(
                doc,
                bullet_lines
            )

            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s+", line):

            flush_paragraph()

            while i < len(lines):

                current = lines[i].strip()

                if not re.match(
                    r"^\d+\.\s+",
                    current
                ):
                    break

                paragraph = doc.add_paragraph(
                    style="List Number"
                )

                paragraph.add_run(
                    re.sub(
                        r"^\d+\.\s*",
                        "",
                        current
                    )
                )

                i += 1

            continue

        paragraph_buffer.append(line)

        i += 1

    flush_paragraph()


def create_document(
    input_file,
    output_file,
    project,
    meeting,
    date_string,
    input_text=None
):

    if input_text is not None:
        markdown = input_text
    else:
        markdown = Path(input_file).read_text(encoding="utf-8")

    doc = Document()

    # --------------------------------------------------------
    # Page setup
    # --------------------------------------------------------

    section = doc.sections[0]

    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # --------------------------------------------------------
    # Default font
    # --------------------------------------------------------

    styles = doc.styles

    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.bold = True

    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.bold = True

    styles["Heading 3"].font.name = "Aptos"
    styles["Heading 3"].font.bold = True

    # --------------------------------------------------------
    # Document header
    # --------------------------------------------------------

    header = section.header

    paragraph = header.paragraphs[0]

    paragraph.text = project

    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.name = "Aptos"

    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    footer = section.footer

    paragraph = footer.paragraphs[0]

    paragraph.text = (
        f"{project}  |  {meeting}  |  {date_string}"
    )

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.name = "Aptos"

    # --------------------------------------------------------
    # Main title
    # --------------------------------------------------------

    title = doc.add_heading(
        "Vergadernotulen",
        level=1
    )

    title.runs[0].font.size = Pt(20)

    # --------------------------------------------------------
    # Meeting information block
    # --------------------------------------------------------

    info_table = doc.add_table(
        rows=3,
        cols=2
    )

    info_table.style = "Table Grid"
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    add_table_borders(info_table)

    info = [
        ("Project", project),
        ("Vergadering", meeting),
        ("Datum", date_string),
    ]

    for row, (label, value) in enumerate(info):

        info_table.cell(row, 0).text = label
        info_table.cell(row, 1).text = value

        set_cell_shading(
            info_table.cell(row, 0),
            "EDEDED"
        )

        set_cell_text_bold(
            info_table.cell(row, 0)
        )

    doc.add_paragraph()

    # --------------------------------------------------------
    # Meeting notes
    # --------------------------------------------------------

    add_markdown_content(
        doc,
        markdown
    )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    output_path = Path(output_file)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    doc.save(output_path)


def main():

    parser = argparse.ArgumentParser()

    parser.add_argument("--input")

    parser.add_argument("--input-text")

    parser.add_argument(
        "--output",
        required=True
    )

    parser.add_argument(
        "--project",
        required=True
    )

    parser.add_argument(
        "--meeting",
        required=True
    )

    parser.add_argument(
        "--date",
        required=True
    )

    args = parser.parse_args()

    if not args.input and args.input_text is None:
        parser.error("one of --input or --input-text is required")

    create_document(
        args.input,
        args.output,
        args.project,
        args.meeting,
        args.date,
        input_text=args.input_text
    )

    print(
        f"Created DOCX: {args.output}"
    )


if __name__ == "__main__":
    main()

