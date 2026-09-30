#!/usr/bin/env python3
"""
Generate Dutch meeting notes from structured WhisperX evidence using Ollama.

Pipeline boundary:
    WhisperX transcript -> extraction/evidence -> this script -> meeting notes

Default model:
    mistral-meeting:16k

Dependencies:
    pip install requests python-docx
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import requests

try:
    from docx import Document
except ImportError:
    Document = None


DEFAULT_MODEL = "mistral-meeting:16k"
DEFAULT_OLLAMA_URL = "http://localhost:11434"


SYSTEM_PROMPT = """Je genereert feitelijke Nederlandse vergaderingsnotities uit
gestructureerde meeting-evidence.

Regels:
1. Gebruik uitsluitend de aangeleverde metadata, deelnemers en evidence.
2. Verzin nooit namen, feiten, deadlines, besluiten, verantwoordelijkheden of data.
3. Gebruik alleen deelnemernamen die expliciet in de participant mapping staan.
4. Maak onderscheid tussen bevestigde besluiten, bevestigde actiepunten,
   voorgestelde acties, vragen, meningen en informatie.
5. Een actie is alleen een bevestigd actiepunt als de evidence expliciet
   ondersteunt dat de actie is afgesproken of toegewezen.
6. Als verantwoordelijke of deadline niet expliciet bekend is, gebruik
   "Niet genoemd" en ga niet raden.
7. Behoud onzekerheid uit de bron.
8. Corrigeer alleen duidelijke transcriptiefouten als de betekenis uit de
   aangeleverde evidence ondubbelzinnig blijkt.
9. Schrijf compact maar informatief.
10. Geef alleen geldig JSON terug, zonder Markdown of code fences.

Gebruik exact deze structuur:
{
  "summary": "string",
  "key_discussion_points": [
    {"topic": "string", "details": "string", "evidence_ids": ["E001"]}
  ],
  "decisions": [
    {"decision": "string", "timestamp": "string", "evidence_ids": ["E001"]}
  ],
  "statements_and_conclusions": [
    {"statement": "string", "speaker": "string", "timestamp": "string", "evidence_ids": ["E001"]}
  ],
  "proposed_actions": [
    {"action": "string", "responsible": "string", "deadline": "string",
     "timestamp": "string", "evidence_ids": ["E001"]}
  ],
  "confirmed_action_items": [
    {"action": "string", "responsible": "string", "deadline": "string",
     "timestamp": "string", "evidence_ids": ["E001"]}
  ],
  "important_dates": [
    {"date": "string", "description": "string", "evidence_ids": ["E001"]}
  ],
  "open_questions": [
    {"question": "string", "evidence_ids": ["E001"]}
  ],
  "additional_notes": ["string"]
}
"""


def load_json(path: Path) -> Any:
    """Load JSON robustly, accepting both UTF-8 and UTF-8-with-BOM files."""
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        raise RuntimeError(f"Bestand niet gevonden: {path}")
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Ongeldige JSON in {path}: {exc}") from exc


def load_evidence(path: Path) -> list[dict[str, Any]]:
    data = load_json(path)
    evidence = data.get("evidence") if isinstance(data, dict) else data
    if not isinstance(evidence, list):
        raise RuntimeError("Evidence JSON moet een 'evidence' lijst bevatten.")
    if not all(isinstance(item, dict) for item in evidence):
        raise RuntimeError("Iedere evidence entry moet een JSON object zijn.")
    return evidence


def load_mapping(path: Path | None, key: str) -> dict[str, str]:
    if path is None:
        return {}
    data = load_json(path)
    if isinstance(data, dict) and isinstance(data.get(key), dict):
        data = data[key]
    if not isinstance(data, dict):
        raise RuntimeError(f"{path} moet een JSON object bevatten.")
    return {str(k): str(v) for k, v in data.items()}


def safe_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "_", value).strip()


def infer_base_name(evidence_path: Path, metadata: dict[str, Any]) -> str:
    name = metadata.get("filename_base") or evidence_path.stem
    return re.sub(r"_evidence$", "", str(name), flags=re.IGNORECASE)


def call_ollama(
    model: str,
    ollama_url: str,
    input_data: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, Any]]:
    payload = {
        "model": model,
        "stream": False,
        "format": "json",
        "options": {"temperature": 0.1},
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": (
                    "Maak de vergaderingsnotities op basis van deze structured data.\n\n"
                    + json.dumps(input_data, ensure_ascii=False, indent=2)
                ),
            },
        ],
    }

    endpoint = ollama_url.rstrip("/") + "/api/chat"

    try:
        response = requests.post(endpoint, json=payload, timeout=1800)
    except requests.RequestException as exc:
        raise RuntimeError(f"Kan Ollama niet bereiken op {endpoint}: {exc}") from exc

    if response.status_code != 200:
        raise RuntimeError(f"Ollama gaf HTTP {response.status_code}: {response.text}")

    try:
        api_result = response.json()
        content = api_result["message"]["content"]
        notes = json.loads(content)
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        raise RuntimeError("Ollama gaf geen geldig JSON-resultaat terug.") from exc

    if not isinstance(notes, dict):
        raise RuntimeError("Het Mistral-resultaat moet een JSON object zijn.")

    return notes, api_result


def as_list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def render_markdown(
    notes: dict[str, Any],
    metadata: dict[str, Any],
    participants: dict[str, str],
) -> str:
    lines = ["# Vergaderingsnotitie", ""]

    for label, key in [
        ("Projectnummer", "project_number"),
        ("Project", "project_name"),
        ("Vergadering", "meeting_details"),
        ("Verwerkingsdatum", "processed_date"),
    ]:
        if metadata.get(key):
            lines.append(f"**{label}:** {metadata[key]}")
    if any(metadata.get(k) for k in ("project_number", "project_name", "meeting_details", "processed_date")):
        lines.append("")

    lines += ["## Deelnemers"]
    if participants:
        lines += [f"- {name}" for name in participants.values()]
    else:
        lines.append("- Niet beschikbaar")
    lines.append("")

    lines += ["## Samenvatting", str(notes.get("summary") or "Niet beschikbaar"), ""]

    lines += ["## Belangrijkste bespreekpunten"]
    points = as_list(notes.get("key_discussion_points"))
    if points:
        for item in points:
            if isinstance(item, dict):
                lines += [f"### {item.get('topic', 'Onderwerp')}",
                          str(item.get("details", "")), ""]
            else:
                lines.append(f"- {item}")
    else:
        lines.append("- Geen expliciete bespreekpunten vastgesteld.")
    lines.append("")

    lines += ["## Besluiten"]
    decisions = as_list(notes.get("decisions"))
    if decisions:
        for item in decisions:
            text = item.get("decision", "") if isinstance(item, dict) else str(item)
            timestamp = item.get("timestamp", "") if isinstance(item, dict) else ""
            lines.append(f"- {text}" + (f" ({timestamp})" if timestamp else ""))
    else:
        lines.append("- Geen expliciete besluiten vastgesteld.")
    lines.append("")

    lines += ["## Uitspraken en conclusies"]
    statements = as_list(notes.get("statements_and_conclusions"))
    if statements:
        for item in statements:
            if isinstance(item, dict):
                speaker = item.get("speaker", "")
                statement = item.get("statement", "")
                timestamp = item.get("timestamp", "")
                prefix = f"**{speaker}:** " if speaker else ""
                suffix = f" ({timestamp})" if timestamp else ""
                lines.append(f"- {prefix}{statement}{suffix}")
            else:
                lines.append(f"- {item}")
    else:
        lines.append("- Geen aanvullende uitspraken of conclusies.")
    lines.append("")

    lines += ["## Voorgestelde acties"]
    proposed = as_list(notes.get("proposed_actions"))
    if proposed:
        for item in proposed:
            if isinstance(item, dict):
                lines.append(
                    f"- {item.get('action', '')} — "
                    f"Verantwoordelijke: {item.get('responsible', 'Niet genoemd')} — "
                    f"Deadline: {item.get('deadline', 'Niet genoemd')}"
                )
            else:
                lines.append(f"- {item}")
    else:
        lines.append("- Geen voorgestelde acties.")
    lines.append("")

    lines += ["## Bevestigde actiepunten"]
    confirmed = as_list(notes.get("confirmed_action_items"))
    if confirmed:
        lines += [
            "| Actie | Verantwoordelijke | Deadline |",
            "|---|---|---|",
        ]
        for item in confirmed:
            if isinstance(item, dict):
                action = str(item.get("action", "")).replace("|", "\\|")
                responsible = str(item.get("responsible", "Niet genoemd")).replace("|", "\\|")
                deadline = str(item.get("deadline", "Niet genoemd")).replace("|", "\\|")
                lines.append(f"| {action} | {responsible} | {deadline} |")
            else:
                lines.append(f"| {item} | Niet genoemd | Niet genoemd |")
    else:
        lines.append("- Geen bevestigde actiepunten.")
    lines.append("")

    lines += ["## Belangrijke data"]
    dates = as_list(notes.get("important_dates"))
    if dates:
        for item in dates:
            if isinstance(item, dict):
                lines.append(f"- **{item.get('date', '')}:** {item.get('description', '')}")
            else:
                lines.append(f"- {item}")
    else:
        lines.append("- Geen belangrijke data vastgesteld.")
    lines.append("")

    lines += ["## Openstaande vragen"]
    questions = as_list(notes.get("open_questions"))
    if questions:
        for item in questions:
            lines.append(f"- {item.get('question', '') if isinstance(item, dict) else item}")
    else:
        lines.append("- Geen openstaande vragen vastgesteld.")
    lines.append("")

    lines += ["## Aanvullende opmerkingen"]
    additional = as_list(notes.get("additional_notes"))
    if additional:
        lines += [f"- {item}" for item in additional]
    else:
        lines.append("- Geen aanvullende opmerkingen.")
    lines.append("")

    return "\n".join(lines)


def write_docx(
    path: Path,
    notes: dict[str, Any],
    metadata: dict[str, Any],
    participants: dict[str, str],
) -> None:
    if Document is None:
        raise RuntimeError("python-docx ontbreekt. Installeer: pip install python-docx")

    doc = Document()
    doc.add_heading("Vergaderingsnotitie", level=0)

    for label, key in [
        ("Projectnummer", "project_number"),
        ("Project", "project_name"),
        ("Vergadering", "meeting_details"),
        ("Verwerkingsdatum", "processed_date"),
    ]:
        if metadata.get(key):
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(str(metadata[key]))

    doc.add_heading("Deelnemers", level=1)
    for name in participants.values():
        doc.add_paragraph(name, style="List Bullet")

    doc.add_heading("Samenvatting", level=1)
    doc.add_paragraph(str(notes.get("summary") or "Niet beschikbaar"))

    doc.add_heading("Belangrijkste bespreekpunten", level=1)
    for item in as_list(notes.get("key_discussion_points")):
        if isinstance(item, dict):
            doc.add_heading(str(item.get("topic", "Onderwerp")), level=2)
            doc.add_paragraph(str(item.get("details", "")))
        else:
            doc.add_paragraph(str(item), style="List Bullet")

    def bullet_section(title: str, items: list[Any], formatter) -> None:
        doc.add_heading(title, level=1)
        if items:
            for item in items:
                doc.add_paragraph(formatter(item), style="List Bullet")
        else:
            doc.add_paragraph("Geen items vastgesteld.")

    bullet_section(
        "Besluiten",
        as_list(notes.get("decisions")),
        lambda x: x.get("decision", "") if isinstance(x, dict) else str(x),
    )

    bullet_section(
        "Uitspraken en conclusies",
        as_list(notes.get("statements_and_conclusions")),
        lambda x: (
            f"{x.get('speaker')}: {x.get('statement')}"
            if isinstance(x, dict) and x.get("speaker")
            else (x.get("statement", "") if isinstance(x, dict) else str(x))
        ),
    )

    bullet_section(
        "Voorgestelde acties",
        as_list(notes.get("proposed_actions")),
        lambda x: (
            f"{x.get('action', '')} — Verantwoordelijke: "
            f"{x.get('responsible', 'Niet genoemd')} — "
            f"Deadline: {x.get('deadline', 'Niet genoemd')}"
            if isinstance(x, dict) else str(x)
        ),
    )

    doc.add_heading("Bevestigde actiepunten", level=1)
    confirmed = as_list(notes.get("confirmed_action_items"))
    if confirmed:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["Actie", "Verantwoordelijke", "Deadline"]):
            cell.text = text
        for item in confirmed:
            row = table.add_row().cells
            if isinstance(item, dict):
                row[0].text = str(item.get("action", ""))
                row[1].text = str(item.get("responsible", "Niet genoemd"))
                row[2].text = str(item.get("deadline", "Niet genoemd"))
            else:
                row[0].text = str(item)
                row[1].text = "Niet genoemd"
                row[2].text = "Niet genoemd"
    else:
        doc.add_paragraph("Geen bevestigde actiepunten.")

    bullet_section(
        "Belangrijke data",
        as_list(notes.get("important_dates")),
        lambda x: (
            f"{x.get('date', '')}: {x.get('description', '')}"
            if isinstance(x, dict) else str(x)
        ),
    )

    bullet_section(
        "Openstaande vragen",
        as_list(notes.get("open_questions")),
        lambda x: x.get("question", "") if isinstance(x, dict) else str(x),
    )

    bullet_section(
        "Aanvullende opmerkingen",
        as_list(notes.get("additional_notes")),
        lambda x: str(x),
    )

    doc.save(path)


def generate_notes(
    evidence_path: Path,
    participants_path: Path | None = None,
    metadata_path: Path | None = None,
    output_dir: Path = Path("3_Meeting Notes"),
    model: str = DEFAULT_MODEL,
    ollama_url: str = DEFAULT_OLLAMA_URL,
    save_raw_response: bool = False,
) -> tuple[Path, Path, Path]:
    evidence = load_evidence(evidence_path)
    participants = load_mapping(participants_path, "participants")
    metadata = load_json(metadata_path) if metadata_path else {}
    if not isinstance(metadata, dict):
        raise RuntimeError("Metadata moet een JSON object zijn.")

    metadata.setdefault("filename_base", infer_base_name(evidence_path, metadata))

    input_data = {
        "meeting_metadata": metadata,
        "participants": participants,
        "evidence": evidence,
    }

    print(f"Evidence entries : {len(evidence)}")
    print(f"Ollama model     : {model}")
    print(f"Evidence         : {evidence_path}")
    print("Mistral wordt aangeroepen...")

    notes, raw_response = call_ollama(model, ollama_url, input_data)

    output_dir.mkdir(parents=True, exist_ok=True)
    base = safe_filename(infer_base_name(evidence_path, metadata))

    json_path = output_dir / f"{base}_meeting_notes.json"
    md_path = output_dir / f"{base}_meeting_notes.md"
    docx_path = output_dir / f"{base}_meeting_notes.docx"

    result = {
        "generated_date": datetime.now().isoformat(timespec="seconds"),
        "model": model,
        "project_number": metadata.get("project_number"),
        "project_name": metadata.get("project_name"),
        "meeting_details": metadata.get("meeting_details"),
        "filename_base": metadata.get("filename_base"),
        "processed_date": metadata.get("processed_date"),
        "participants": participants,
        "notes": notes,
    }
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    md_path.write_text(render_markdown(notes, metadata, participants), encoding="utf-8")
    write_docx(docx_path, notes, metadata, participants)

    if save_raw_response:
        raw_path = output_dir / f"{base}_ollama_response.json"
        raw_path.write_text(json.dumps(raw_response, ensure_ascii=False, indent=2), encoding="utf-8")

    print("Klaar.")
    print(f"JSON : {json_path}")
    print(f"MD   : {md_path}")
    print(f"DOCX : {docx_path}")
    return json_path, md_path, docx_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate meeting notes from WhisperX evidence via Ollama.")
    parser.add_argument("--evidence", required=True, type=Path)
    parser.add_argument("--participants", type=Path)
    parser.add_argument("--metadata", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("3_Meeting Notes"))
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--ollama-url", default=DEFAULT_OLLAMA_URL)
    parser.add_argument("--save-raw-response", action="store_true")
    args = parser.parse_args()

    if not args.evidence.exists():
        print(f"ERROR: evidence-bestand bestaat niet: {args.evidence}", file=sys.stderr)
        return 1

    try:
        generate_notes(
            evidence_path=args.evidence,
            participants_path=args.participants,
            metadata_path=args.metadata,
            output_dir=args.output_dir,
            model=args.model,
            ollama_url=args.ollama_url,
            save_raw_response=args.save_raw_response,
        )
    except RuntimeError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
